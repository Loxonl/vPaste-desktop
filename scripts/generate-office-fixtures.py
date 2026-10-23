"""Rebuild public, synthetic Office fixtures (authoring only; not an app dependency).

Requires python-docx, openpyxl and Pillow. The generated documents contain no
user clipboard data, filesystem paths, comments, macros or external resources.
"""

import argparse
import io
import json
import random
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree

from docx import Document
from docx.shared import Inches, RGBColor
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as SheetImage
from openpyxl.styles import Border, Font, PatternFill, Side
from PIL import Image, ImageDraw


ROOT = Path(__file__).resolve().parents[1]
CATALOG = ROOT / "tests/fixtures/office/catalog.json"
FIXED_TIME = datetime(2000, 1, 1, tzinfo=timezone.utc)


def stable_archive(path: Path) -> None:
    """Remove ZIP timestamps/host attributes so regeneration stays reviewable."""
    with zipfile.ZipFile(path) as source:
        entries = {
            name: source.read(name) for name in source.namelist()
            if not name.startswith("customXml/") and name != "docProps/thumbnail.jpeg"
        }
    for name, data in list(entries.items()):
        if name.endswith(".rels") or name == "[Content_Types].xml":
            root = ElementTree.fromstring(data)
            for element in list(root):
                target = element.get("Target", element.get("PartName", ""))
                if "customXml/" in target or target.endswith("thumbnail.jpeg"):
                    root.remove(element)
            entries[name] = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
        if name == "docProps/core.xml":
            entries[name] = re.sub(
                rb"(<dcterms:modified\b[^>]*>)[^<]+",
                rb"\g<1>2000-01-01T00:00:00Z", data,
            )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for name, data in sorted(entries.items()):
            entry = zipfile.ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.create_system = 0
            entry.external_attr = 0
            target.writestr(entry, data)


def generate(destination: Path) -> None:
    catalog = json.loads(CATALOG.read_text(encoding="utf-8"))
    destination.mkdir(parents=True, exist_ok=True)
    for image in catalog["images"]:
        size = (image["width"], image["height"])
        if image["key"] == "largePicture":
            pixels = random.Random(42).randbytes(size[0] * size[1] * 3)
            picture = Image.frombytes("RGB", size, pixels)
        else:
            color = {"pictureA": "#234e9b", "pictureB": "#31a56b", "icon": "#d96935"}[image["key"]]
            picture = Image.new("RGB", size, color)
            ImageDraw.Draw(picture).rectangle((size[0] // 4, size[1] // 4, size[0] // 2, size[1] // 2), fill="white")
        picture.save(destination / image["file"], format="PNG", optimize=False)

    document = Document()
    document.core_properties.author = "vPaste synthetic fixtures"
    document.core_properties.last_modified_by = "vPaste synthetic fixtures"
    document.core_properties.created = FIXED_TIME
    document.core_properties.modified = FIXED_TIME
    document.core_properties.title = catalog["marker"]
    document.core_properties.comments = "Generated public test content."
    for case in catalog["cases"]:
        if case["group"] != "word":
            continue
        document.add_heading(f'{case["code"]} - {case["name"]}', level=1)
        if case["id"] == "office-word-formatted":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(catalog["marker"])
            run.bold = True
            run.font.color.rgb = RGBColor.from_string("234E9B")
            run = paragraph.add_run(" - Word sample")
            run.font.highlight_color = 7  # yellow
            continue
        keys = {
            "office-word-image-row": ["pictureA"],
            "office-word-mixed": ["pictureA", "pictureB", "icon"],
            "office-word-large-picture": ["largePicture"],
            "office-word-consecutive": ["pictureA", "pictureB"],
        }[case["id"]]
        if case["plainText"]:
            document.add_paragraph(case["plainText"])
        for key in keys:
            image = next(item for item in catalog["images"] if item["key"] == key)
            document.add_paragraph().add_run().add_picture(
                str(destination / image["file"]), width=Inches(image["width"] / 96),
            )
    word_path = destination / "vPaste-Office.docx"
    document.save(word_path)
    stable_archive(word_path)

    workbook = Workbook()
    workbook.properties.creator = "vPaste synthetic fixtures"
    workbook.properties.lastModifiedBy = "vPaste synthetic fixtures"
    workbook.properties.created = FIXED_TIME
    workbook.properties.modified = FIXED_TIME
    workbook.properties.title = catalog["marker"]
    sheet = workbook.active
    sheet.title = "Samples"
    for row in [("Product", "Units", "Region"), ("Sample A", 12, "North"), ("Sample B", 24, "South"), ("Sample C", 36, "East")]:
        sheet.append(row)
    for row in sheet["A1:C4"]:
        for cell in row:
            cell.border = Border(bottom=Side(style="thin", color="234E9B"))
    for cell in sheet[1][:3]:
        cell.font = Font(bold=True, color="234E9B")
        cell.fill = PatternFill("solid", fgColor="DCE6F1")
    sheet["B3"].fill = PatternFill("solid", fgColor="FFF59D")
    sheet.merge_cells("A6:C6")
    sheet["A6"] = "Synthetic sales"
    for column in "ABC":
        sheet.column_dimensions[column].width = 18
    chart = BarChart()
    chart.title = "Synthetic units"
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=4), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=4))
    sheet.add_chart(chart, "E5")
    sheet.add_image(SheetImage(io.BytesIO((destination / "picture-a.png").read_bytes())), "A8")
    guide = workbook.create_sheet("Guide")
    guide.append(("Case", "Check", "Steps"))
    for case in catalog["cases"]:
        if case["group"] == "excel":
            guide.append((case["code"], case["name"], " ".join(case["manualSteps"])))
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 34
    guide.column_dimensions["C"].width = 100
    excel_path = destination / "vPaste-Office.xlsx"
    workbook.save(excel_path)
    stable_archive(excel_path)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=CATALOG.parent)
    arguments = parser.parse_args()
    generate(arguments.output.resolve())
    print("Generated synthetic Office fixtures.")
